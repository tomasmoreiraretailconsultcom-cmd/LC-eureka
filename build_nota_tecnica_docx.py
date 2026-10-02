import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_document():
    doc = Document()

    # Define color palette
    COLOR_PRIMARY = RGBColor(15, 76, 129)      # Classic Navy / #0F4C81
    COLOR_SECONDARY = RGBColor(30, 41, 59)     # Slate Navy / #1E293B
    COLOR_ACCENT = RGBColor(225, 29, 72)       # Crimson / #E11D48
    COLOR_MUTED = RGBColor(100, 116, 139)      # Slate Muted / #64748B
    COLOR_TEXT = RGBColor(51, 65, 85)          # Body Text / #334155

    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base style configurations
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Helper function for cell background color
    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    # Helper function for cell margins/padding
    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._element.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    # Helper function to create a callout box
    def add_callout(text, title="NOTA DE DESTAQUE", fill_hex="F8FAFC", border_color="0F4C81"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
        
        tcPr = cell._element.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        run_t = p.add_run(f"📌 {title}\n")
        run_t.bold = True
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = COLOR_PRIMARY
        
        run_c = p.add_run(text)
        run_c.font.name = 'Calibri'
        run_c.font.size = Pt(10)
        run_c.font.color.rgb = COLOR_SECONDARY
        
        doc.add_paragraph()

    # Helper for Formula Box
    def add_formula_box(equation_lines, title="Fórmula Matemática"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._element.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="0F4C81"/><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        
        if title:
            r_title = p.add_run(f"{title}\n")
            r_title.bold = True
            r_title.font.name = 'Calibri'
            r_title.font.size = Pt(10)
            r_title.font.color.rgb = COLOR_PRIMARY

        for idx, line in enumerate(equation_lines):
            r = p.add_run(line + ("\n" if idx < len(equation_lines)-1 else ""))
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_SECONDARY
            r.bold = True
            
        doc.add_paragraph()

    # ----------------------------------------------------
    # DOCUMENT HEADER / TITLE
    # ----------------------------------------------------
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(0)
    p_pre.paragraph_format.space_after = Pt(2)
    r_pre = p_pre.add_run("LIFECYCLE INTELLIGENCE ENGINE  |  NOTA TÉCNICA DE MODELAÇÃO")
    r_pre.font.name = 'Calibri'
    r_pre.font.size = Pt(9.5)
    r_pre.font.bold = True
    r_pre.font.color.rgb = COLOR_MUTED

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("Sincronização Matemática da Curva de Qualidade com o Ponto B (Apodrecimento / Fim de Vida)")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(0)
    p_meta.paragraph_format.space_after = Pt(16)
    r_meta = p_meta.add_run("Data: Outubro 2026   |   Versão do Modelo: LC_model_v0_2_0   |   Módulo: Curva de Degradação & Shelf-Life")
    r_meta.font.name = 'Calibri'
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = COLOR_MUTED
    r_meta.italic = True

    # ----------------------------------------------------
    # SECTION 1: CONTEXTO E PROBLEMA
    # ----------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Diagnóstico do Problema Observado")
    r_h1.font.name = 'Calibri'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "Nas simulações visuais da UI, observou-se uma assincronia entre o marcador visual do "
        "Ponto B (Losango Vermelho — Apodrecimento / Incomestibilidade Biológica) e o declínio "
        "da linha contínua azul (Índice Global de Qualidade Comercial, em %)."
    )

    doc.add_paragraph(
        "Especificamente, quando o fungo atingia o limiar crítico de condenação biológica "
        "(ex.: no Dia 7 com Mold = 50%), o gráfico plotava o marcador de Apodrecimento no Dia 7, "
        "mas a curva de qualidade permanecia em cerca de ~34%, descendo gradualmente até aos 0% "
        "apenas no Dia 11."
    )

    add_callout(
        "Incongruência Perceptiva: Se o fruto atingiu o Ponto B (está apodrecido e impróprio para consumo), "
        "a sua qualidade comercial e organolética remanescente deve ser nula (0.0%). Uma peça com bolor invasivo "
        "não possui 34% de qualidade residual aproveitável.",
        title="INCONGRUÊNCIA FÍSICO-BIOLÓGICA",
        border_color="E11D48"
    )

    # ----------------------------------------------------
    # SECTION 2: FORMULAÇÃO ATUAL (ANTES)
    # ----------------------------------------------------
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. Formulação Matemática Atual (Antes da Correção)")
    r_h2.font.name = 'Calibri'
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "No ficheiro LC_model_v0_2_0.py, o cálculo da qualidade em cada instante t era definido "
        "através da ponderação dos parâmetros físico-químicos e de um fator de penalização de bolor aditivo:"
    )

    add_formula_box([
        "Qualidade_bruta(t) = w_firmeza * (Firmeza(t)/Firmeza_0) + w_brix * (Brix(t)/Brix_0) + w_acidez * (Acidez(t)/Acidez_0)",
        "",
        "Penalização_Bolor(t) = 1.0 - (mold_max_penalty * Mold(t))",
        "onde mold_max_penalty = 0.65",
        "",
        "Qualidade_final(t) = Qualidade_bruta(t) * Penalização_Bolor(t) * cap_decay(t)"
    ], title="Fórmulas Atuais (Desacopladas)")

    doc.add_paragraph(
        "Causa raiz da discrepância matemática:\n"
        "• No momento do Ponto B, Mold(t) = 0.50 (50%).\n"
        "• A penalização aplicada era apenas: 1.0 - (0.65 * 0.50) = 1.0 - 0.325 = 0.675 (67.5%).\n"
        "• Como a polpa e o açúcar ainda mantinham ~50% de atributos, Qualidade(t) = 50% * 0.675 = 33.7%.\n"
        "• O Ponto B era calculado isoladamente pela condição Mold(t) >= 0.50, enquanto a curva continuava viva."
    )

    # ----------------------------------------------------
    # SECTION 3: NOVA FORMULAÇÃO INTEGRADA (DEPOIS)
    # ----------------------------------------------------
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. Nova Formulação Matemática Unificada (Proposta)")
    r_h3.font.name = 'Calibri'
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "Para assegurar sincronismo total e rigor fisiológico, a qualidade global passa a ser o produto "
        "da qualidade organolética pelo Fator de Sanidade Biológica e pelo Fator de Senescência Fisiológica."
    )

    add_formula_box([
        "1. Fator de Sanidade Biológica (Bolor / Apodrecimento):",
        "   F_sanidade(t) = max(0.0,  1.0 - (Mold(t) / 0.50))",
        "",
        "2. Fator de Senescência Fisiológica (Cap de Shelf-Life Física):",
        "   F_senescencia(t) = clip( remaining_shelf_life_fisica(t) / max(1.0, 0.10 * SL_referencia),  0.0,  1.0 )",
        "",
        "3. Índice Global de Qualidade (%):",
        "   Qualidade(t) = Qualidade_organoletica(t) * F_sanidade(t) * F_senescencia(t) * 100",
        "",
        "4. Definição Unificada do Ponto B (Intercepto Fim de Vida):",
        "   Ponto_B = min { t >= 0 : Qualidade(t) <= 0.01% }"
    ], title="Fórmulas Propostas (Sincronização Perfeita)")

    doc.add_paragraph(
        "Comportamento e Propriedades da Nova Formulação:\n"
        "1. Enquanto Mold(t) = 0.0%, F_sanidade = 1.0 (qualidade segue a evolução normal de maturação/açúcares).\n"
        "2. Quando Mold(t) progride de 0% a 50%, F_sanidade desce suavemente de 1.0 a 0.0.\n"
        "3. No instante exato em que Mold(t) atinge 50%, F_sanidade = 0.0, forçando Qualidade(t) = 0.0%.\n"
        "4. O marcador Ponto B passa a ser exatamente o ponto onde a linha azul toca o eixo horizontal (0%)."
    )

    # ----------------------------------------------------
    # SECTION 4: TABELA COMPARATIVA ANTES VS DEPOIS
    # ----------------------------------------------------
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. Quadro Comparativo: Antes vs. Depois")
    r_h4.font.name = 'Calibri'
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = COLOR_PRIMARY

    # Create Table
    tbl = doc.add_table(rows=5, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    headers = ["Aspeto / Métrica", "Modelo Atual (Desacoplado)", "Modelo Proposto (Sincronizado)"]
    widths = [Inches(1.8), Inches(2.3), Inches(2.4)]

    # Header row
    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.width = widths[col_idx]
        set_cell_background(cell, "0F4C81")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("Penalização de Bolor", "Linear atenuada: 1 - 0.65 * mold\n(nunca anula a qualidade)", "Função de extinção: max(0, 1 - mold / 0.50)\n(anula qualidade aos 50% de bolor)"),
        ("Valor da Linha Azul no Ponto B", "Permanece entre 30% e 40%\n(cria confusão visual)", "Atinge rigorosamente 0.0%\n(convergência total)"),
        ("Definição do Ponto B", "Independente da linha de qualidade\n(condição externa mold >= 0.5)", "Sincronizada: dia em que Qualidade(t) == 0%\n(integridade matemática)"),
        ("Interpretação Comercial", "Fruto com bolor crítico ainda mostra 'qualidade'", "Fruto com bolor crítico tem 0% de valor comercial")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = tbl.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = COLOR_PRIMARY
            elif col_idx == 1:
                run.font.color.rgb = COLOR_MUTED
            else:
                run.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph()

    # ----------------------------------------------------
    # SECTION 5: IMPACTO E PRÓXIMOS PASSOS
    # ----------------------------------------------------
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    r_h5 = h5.add_run("5. Impacto na Simulação e Próximos Passos")
    r_h5.font.name = 'Calibri'
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "Com esta alteração matemática:\n"
        "1. Os dois motores de cálculo (Simulação Luís Paulo e Simulação Sofia Machado) mantêm as suas taxas diferenciais de degradação físico-química, mas respeitam o mesmo princípio biológico de terminação.\n"
        "2. A curva na UI Streamlit desce suave e realisticamente até ao ponto exato onde se encontra o marcador vermelho (Ponto B), eliminando caudas residuais irreais.\n"
        "3. Os 65 cenários Excel e a UI passam a apresentar consistência a 100%."
    )

    add_callout(
        "Para aplicar estas fórmulas no motor de cálculo, basta atualizar as linhas de agregação final "
        "das funções 'run_simulation_prof_luis_paulo' e 'run_simulation_sofia_machado' no ficheiro LC_model_v0_2_0.py.",
        title="APLICAÇÃO EM CÓDIGO",
        fill_hex="F0FDF4",
        border_color="16A34A"
    )

    output_path = r"c:\Lifecycle\LC-eureka\Nota_Tecnica_Sincronizacao_Ponto_B_Qualidade.docx"
    doc.save(output_path)
    print(f"Documento criado com sucesso em: {output_path}")

if __name__ == "__main__":
    create_document()
