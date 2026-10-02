import openpyxl
import pandas as pd
import datetime
import os
import re
import unicodedata

def clean(s):
    if s is None:
        return ''
    s = str(s)
    return unicodedata.normalize('NFKD', s).encode('ASCII', 'ignore').decode('ASCII')

def slugify(s):
    s = clean(s).lower()
    s = re.sub(r'[^a-z0-9]+', '_', s).strip('_')
    return s

wb = openpyxl.load_workbook(r'c:\Lifecycle\LC-eureka\Cenarios LC.xlsx', data_only=True)
ws = wb.active

out_dir = r'c:\Lifecycle\LC-eureka\cenarios_excel'
os.makedirs(out_dir, exist_ok=True)

manifest = []
base_date = datetime.date.today() - datetime.timedelta(days=7)

count = 0
for r in range(6, ws.max_row + 1):
    fruit_raw = clean(ws.cell(r, 3).value)
    scenario_raw = clean(ws.cell(r, 4).value)
    expected_life = clean(ws.cell(r, 11).value)
    pkg_raw = clean(ws.cell(r, 14).value) or 'bulk'
    
    if not fruit_raw or not scenario_raw:
        continue
    
    count += 1
    
    # Extract Temp, RH, Ethylene
    temp_match = re.search(r'(-?\d+(?:\.\d+)?)\s*C', scenario_raw)
    rh_match = re.search(r'(\d+(?:\.\d+)?)\s*%\s*HR', scenario_raw)
    eth_match = re.search(r'(\d+(?:\.\d+)?)\s*ppm', scenario_raw, re.IGNORECASE)
    
    temp = float(temp_match.group(1)) if temp_match else 20.0
    rh = float(rh_match.group(1)) if rh_match else 85.0
    eth = float(eth_match.group(1)) if eth_match else None
    
    # Generate clean filename
    fruit_slug = slugify(fruit_raw)
    scen_short = slugify(scenario_raw)[:32].rstrip('_')
    filename = f'{count:02d}_{fruit_slug}_{scen_short}.xlsx'
    filepath = os.path.join(out_dir, filename)
    
    # Create 7 days of daily readings with exact scenario conditions
    days_hist = 7
    dates = [(base_date + datetime.timedelta(days=i)).strftime('%Y-%m-%d') for i in range(days_hist)]
    
    data = {
        'Segment_ID': [1] * days_hist,
        'Date': dates,
        'Temperature_C': [temp] * days_hist,
        'Humidity_Percent': [rh] * days_hist,
        'Region': ['PT-LVT'] * days_hist,
        'Packaging': [pkg_raw] * days_hist
    }
    
    if eth is not None and eth > 0:
        data['Ethylene_ppm'] = [eth] * days_hist
        
    df = pd.DataFrame(data)
    df.to_excel(filepath, index=False)
    
    manifest.append({
        'num': count,
        'row': r,
        'fruit': fruit_raw,
        'scenario': scenario_raw,
        'temp': temp,
        'rh': rh,
        'eth': eth if eth else '-',
        'pkg': pkg_raw,
        'expected_life': expected_life,
        'filename': filename
    })

print(f'Total de {count} ficheiros Excel gerados na pasta {out_dir}!')

# Save a markdown manifest
md_path = os.path.join(out_dir, 'README_CENARIOS.md')
with open(md_path, 'w', encoding='utf-8') as f_md:
    f_md.write('# 📁 Índice de Ficheiros Excel de Cenários para Teste (Lifecycle)\n\n')
    f_md.write('| # | Fruto | Cenário | Temp (°C) | HR (%) | Etileno | Embalagem | Vida Esperada | Ficheiro Excel |\n')
    f_md.write('|:---:|:---|:---|:---:|:---:|:---:|:---:|:---|:---|\n')
    for m in manifest:
        f_md.write(f"| {m['num']:02d} | {m['fruit']} | {m['scenario']} | {m['temp']}°C | {m['rh']}% | {m['eth']} | `{m['pkg']}` | {m['expected_life']} | [`{m['filename']}`](./{m['filename']}) |\n")

print('Manifesto README_CENARIOS.md criado com sucesso!')
